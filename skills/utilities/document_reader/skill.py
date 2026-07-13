"""Document Reader Skill — Read PDF, DOCX, TXT files."""
import logging
import os
import re
from typing import Dict, Any
from skills import BaseSkill
logger = logging.getLogger(__name__)

class DocumentReaderSkill(BaseSkill):
    def __init__(self):
        super().__init__(name="document_reader", category="utilities", capabilities=["document_reader"])

    def execute(self, capability: str, params: Dict[str, Any]) -> Dict[str, Any]:
        method = params.get("method", "read")
        if method == "read": return self.read(params)
        elif method == "metadata": return self.get_metadata(params)
        return {"error": f"Unknown method: {method}"}

    def read(self, params: Dict[str, Any]) -> Dict[str, Any]:
        path = params.get("path", "")
        if not path or not os.path.exists(path): return {"success": False, "error": f"File not found: {path}"}
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".pdf":
                try:
                    import PyPDF2
                    with open(path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        text = "\n".join(page.extract_text() or "" for page in reader.pages)
                    return {"success": True, "format": "pdf", "pages": len(reader.pages), "text": text[:5000], "path": path}
                except ImportError:
                    return {"success": False, "error": "PyPDF2 required. Install: pip install PyPDF2"}
            elif ext in (".docx", ".doc"):
                try:
                    import docx
                    doc = docx.Document(path)
                    text = "\n".join(p.text for p in doc.paragraphs)
                    return {"success": True, "format": "docx", "paragraphs": len(doc.paragraphs), "text": text[:5000], "path": path}
                except ImportError:
                    return {"success": False, "error": "python-docx required. Install: pip install python-docx"}
            elif ext in (".txt", ".md", ".csv", ".json", ".py", ".js", ".html"):
                with open(path, 'r', encoding='utf-8', errors='replace') as f:
                    text = f.read()
                return {"success": True, "format": ext.lstrip("."), "text": text[:5000], "path": path}
            else:
                return {"success": False, "error": f"Unsupported format: {ext}"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def get_metadata(self, params: Dict[str, Any]) -> Dict[str, Any]:
        path = params.get("path", "")
        if not os.path.exists(path): return {"success": False, "error": "File not found"}
        stat = os.stat(path)
        return {"success": True, "path": path, "size_bytes": stat.st_size, "modified": stat.st_mtime, "extension": os.path.splitext(path)[1]}
